"""Build the final deck, report and timed script from preserved measured receipts."""
import json
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / 'docs'
SLIDES_URL = 'https://docs.google.com/presentation/d/1J2m_cpXyvb-5_Xabs2DsD3EnEV6B6GicF5OP2s5FohU/edit'
REPO_URL = 'https://github.com/dragonchan95/Semiconductor-AI-Agent'
NAMES = {'logistic':'Logistic Regression','decision_tree':'Decision Tree',
         'random_forest':'Random Forest','gradient_boosting':'Gradient Boosting',
         'xgboost':'XGBoost','lightgbm':'LightGBM','catboost':'CatBoost','knn':'KNN',
         'svm':'SVM (RBF)','mlp':'MLP / ANN','adaboost':'AdaBoost',
         'hist_gradient':'HistGradientBoosting','extra_trees':'ExtraTrees',
         'gaussian_nb':'GaussianNB','resnet_small':'축소 ResNet','vgg_small':'축소 VGGNet',
         'autoencoder':'Autoencoder'}
TEAL='147D82'; INK='173A49'; MUTED='617681'; BG='F6F8FA'; LINE='D8E3E7'; PALE='E8F3F3'; ORANGE='FFF0DD'


def load(name):
    return json.loads((DOCS / 'evaluation' / f'{name}.json').read_text(encoding='utf-8'))


def pct(value):
    return str(Decimal(str(value * 100)).quantize(Decimal('0.1'), rounding=ROUND_HALF_UP)) + '%'


def ranking(trials, policy):
    def key(t):
        m=t['validation']
        return (m['recall']>=policy['min_recall'] and m['fpr']<=policy['max_fpr'],
                -max(0,policy['min_recall']-m['recall'])-max(0,m['fpr']-policy['max_fpr']),
                m['recall'],-m['fpr'],t['id'])
    complete=[t for t in trials if t['status']=='completed']
    count=max(len(t['products']) for t in complete)
    best={}
    for t in complete:
        if len(t['products'])==count and (t['model'] not in best or key(t)>key(best[t['model']])):
            best[t['model']]=t
    return sorted(best.values(),key=key,reverse=True)


def make_content():
    receipts={n:load(n) for n in ['synthetic_process','secom','images']}
    r=receipts['synthetic_process']['result']; w=r['winner']; b=r['baseline_test']; t=r['test']
    policy=receipts['synthetic_process']['plan']['policy']; ranks=ranking(r['trials'],policy)
    mg=load('model_management')['comparison']; retrain=mg['same_model_retraining']
    trace=load('langsmith_check') if (DOCS/'evaluation/langsmith_check.json').exists() else None
    windows=load('windows_check') if (DOCS/'evaluation/windows_check.json').exists() else None
    test_state='30개 테스트 통과' if windows and windows.get('tests_passed') else '기존 30개 테스트 · Windows 재검증 기록 별도'
    content=[]
    def add(title,context,body,seconds=0,note='',table=None,image=None,kind=None,sources=None):
        content.append(dict(title=title,context=context,body=body,seconds=seconds,note=note,
                            table=table,image=image,kind=kind,sources=sources or []))
    add('반도체 검사 모델\n실험 도우미','AI AGENT MINI PROJECT · 과제 2',[
        '데이터를 점검하고, 승인받아 실험하고, 다음 비교를 제안합니다.',
        'SKALA 광주 2반 · G045 김용찬 · 2026.10.22'],15,
        '안녕하세요. SKALA 광주 2반 G045 김용찬입니다. 반도체 검사 데이터를 넣으면 모델을 비교하고, 결과에 따라 다음 실험을 제안하는 Agent를 만들었습니다.',kind='cover')
    add('데이터가 다르면 비교할 모델도 달라집니다','문제 · 반복 실험을 연결하는 도구가 필요합니다',[
        '사용자: 새 검사 데이터로 정상·불량 분류 모델을 만드는 엔지니어',
        '전처리 → 모델 비교 → 근거 정리 → 추가 실험을 매번 반복합니다.'],25,
        '스칼라에서 배운 모델을 바탕으로 반도체 자료에서 참고한 후보도 함께 넣었습니다. 검사 이미지, 로그, 시계열처럼 데이터 형태에 따라 사용하는 모델이 달랐습니다. 이번에는 숫자 센서와 검사 이미지의 정상·불량 분류로 범위를 줄였습니다. 모델 하나를 추천하는 대신, 엔지니어가 반복하는 비교와 기록 과정을 연결하는 것이 목적입니다.',
        table=[['데이터','교육자료의 모델 예시','이번 구현'],
               ['숫자 센서','Tree · KNN · SVM · ANN','여러 분류 모델 실제 학습'],
               ['검사 이미지','ResNet · VGGNet · Autoencoder','자체 축소 구조 3종 실제 학습'],
               ['로그 / 시계열','BERT·GPT / LSTM·ARIMA','향후 확장']],
        sources=['docs/CURRICULUM_MODELS.md','반도체 교육 사진의 분류를 표로 재구성'])
    add('Agent는 계획과 근거, 엔지니어는 승인을 맡습니다','아키텍처 · 실제 실행 순서',[
        'LangGraph로 점검·검색·계획 역할을 연결하고, 승인된 계획만 실행합니다.',
        '목표 미달이면 새 모델 탐색 / 기존 모델 점검으로 이어갑니다.'],35,
        '구조입니다. 데이터 점검, 근거 검색, 실험 계획을 LangGraph로 연결했습니다. 계획과 데이터의 해시를 저장하고 엔지니어가 승인한 뒤에만 학습합니다. 결과 검토에서는 부족한 지표를 보여 주고 새 모델 비교나 기존 모델 점검으로 이어집니다. 학습과 수치는 실험 엔진이 계산하고, LLM은 제한된 근거와 설명을 보조합니다. 자율 Supervisor나 자동 배포까지 구현한 것은 아닙니다.',kind='architecture',
        sources=['app/planner.py','app/workbench.py','app/store.py','mcp_server.py'])
    add('① 데이터를 선택하고 비교 계획을 요청합니다','같은 합성 센서 사례를 끝까지 따라갑니다',[
        '제품 A–D · 1,280건 · 불량 201건',
        'A–D는 가상 제품명입니다.',
        '첫 사용: 예제 데이터 준비 → 센서 선택 → 비교 계획 제안받기'],20,
        '한 사례를 끝까지 보겠습니다. 제품 A부터 D의 합성 센서 1,280건을 선택합니다. 불량 정답은 201건이고 제품 이름은 가상입니다. 처음 실행하면 예제 데이터를 준비하고, 센서를 선택한 뒤 비교 계획을 요청합니다.',image='app-data.jpg',
        sources=['app/data.py','docs/evaluation/synthetic_process.json'])
    add('② 전처리·모델·분할을 확인한 뒤 승인합니다','승인 전에는 학습하지 않습니다',[
        'Train 768 / Validation 256 / Test 256',
        '같은 생산 묶음(lot)은 구간 사이에 섞지 않습니다.',
        '여러 후보·불균형 설정을 비교한 뒤 상위 후보의 제품 조합 확장',
        '최대 52조건 · 완전탐색은 아닙니다.'],25,
        'Agent가 결측과 불량 비율을 확인하고 전처리와 모델 계획을 제안합니다. 수업 기준 6대2대2를 사전에 고정하고 같은 생산 묶음이 평가에 섞이지 않도록 나눴습니다. 스칼라에서 배운 모델과 반도체 자료에서 참고한 후보를 함께 비교하고 상위 후보의 제품 조합을 확장합니다. 이 계획을 승인해야 실제 학습합니다.',image='app-plan.jpg',
        sources=['app/planner.py','app/experiments.py'])
    show=[ranks[0],ranks[1],next(x for x in ranks if x['model']=='adaboost'),next(x for x in ranks if x['model']=='logistic')]
    table=[['전체 제품 학습 · Validation','Recall','정상 오판률','F1']]+[
        [NAMES[x['model']],pct(x['validation']['recall']),pct(x['validation']['fpr']),pct(x['validation']['f1'])] for x in show]
    add('③ Validation으로 후보를 고르고 이유를 남깁니다','Test 결과를 보고 모델 순위를 바꾸지 않습니다',[
        '선택 규칙: Recall ≥ 90% · 정상 오판률 ≤ 15% 목표 충족 우선',
        '충족 후보 중 Recall 우선, 동률이면 정상 오판률이 낮은 후보',
        '상위 2모델의 제품 조합 비교 → CatBoost · A+B+C 선택'],35,
        '왜 CatBoost를 골랐는지입니다. 먼저 Validation에서 두 목표를 충족했는지 확인하고 Recall을 우선합니다. 동률이면 정상 오판률이 낮은 후보가 앞섭니다. AdaBoost는 F1이 더 높지만 Recall 우선 규칙에서는 CatBoost가 앞섰습니다. 이후 상위 두 알고리즘의 제품 조합을 비교해 A B C로 학습한 CatBoost를 골랐습니다. 이 규칙은 실험 전에 정한 교육용 목표입니다.',table=table,
        sources=['docs/evaluation/synthetic_process.json','app/metrics.py:selection_key'])
    add('④ 더 찾았지만, 목표에는 못 미쳤습니다','같은 Test 256건 · 불량 32건 / 정상 224건',[
        f'기본: Logistic Regression / 선택 후보: {NAMES[w["model"]]} · {"+".join(w["products"])}',
        f'후보 Recall {pct(t["recall"])} < 목표 90% → 추가 실험 필요'],35,
        f'마지막 Test입니다. 같은 불량 {t["positive_count"]}개 중 기본 Logistic은 {b["tp"]}개, 후보 CatBoost는 {t["tp"]}개를 찾았습니다. 대신 정상을 불량으로 오판한 수는 {b["fp"]}개에서 {t["fp"]}개로 늘었습니다. 후보 Recall {pct(t["recall"])}는 목표 90퍼센트에 못 미쳤습니다. 선택 후보와 현장 적용 확정은 다른 단계이며, 이 결과를 보고 다음 실험으로 이어갑니다.',
        table=[['같은 Test에서 비교','기본 모델','선택 후보'],
               ['불량 발견',f'{b["tp"]} / 32개',f'{t["tp"]} / 32개'],
               ['놓친 불량',f'{b["fn"]}개',f'{t["fn"]}개'],
               ['정상 오판',f'{b["fp"]} / 224개',f'{t["fp"]} / 224개'],
               ['Recall',pct(b['recall']),pct(t['recall'])]],
        sources=['docs/evaluation/synthetic_process.json'])
    add('⑤ 공식 문서에서 새 후보를 찾아 다시 비교합니다','다음 실험 · 모델 탐색 RAG',[
        '공식 모델 목록·문서 조회 → 근거 검색 → 새 비교 계획 → 승인·학습',
        'ExtraTrees · GaussianNB를 기존 후보와 비교하고 목록을 갱신',
        '기록된 추가 실험에서는 CatBoost 유지',
        '현재는 검토된 두 adapter 범위 · 임의 코드 자동 실행은 미구현'],35,
        '다음 행동은 새 모델 비교입니다. 공식 모델 목록과 문서를 조회하고, RAG로 근거를 연결해 ExtraTrees와 Gaussian Naive Bayes를 제안합니다. 승인받아 기존 후보와 같은 조건으로 학습하고 후보 목록에 기록했습니다. 이 추가 실험에서는 CatBoost를 유지했습니다. 현재는 검토된 두 모델 연결을 지원합니다. 기존 Test를 재사용한 탐색 비교여서 최종 적용에는 새로운 독립 데이터가 필요합니다.',image='app-discovery.jpg',
        sources=['app/discovery.py','docs/evaluation/model_discovery.json'])
    add('실험 경로와 실패 상태를 기록합니다','적용 기술 · 관측 가능성 및 평가',[
        'LangGraph 역할 분리 · MCP 4도구와 지식 Resource · 의존성 주입',
        'RAG: 교육 근거 14개 + 공식 외부 문서 검색',
        ('LangSmith APAC: 실제 계획 생성과 하위 실행 API 재조회 성공' if trace else 'LangSmith: 연결 검증 결과 별도 확인'),
        f'{test_state} · LLM 응답 실패/미사용은 규칙 설명으로 표시'],30,
        '기술은 LangGraph 역할 분리, MCP 네 도구와 지식 리소스, RAG와 의존성 주입을 적용했습니다. LangSmith에는 이 노트북에서 실제 계획 생성과 하위 실행 기록을 보내고 다시 조회했습니다. 다만 이 노트북의 현재 모드는 로컬 LLM 미사용입니다. 기존 노트북의 LLM 실행 기록과 구분하고, 응답 실패는 규칙 설명으로 표시합니다. LLM-as-a-Judge는 미구현입니다.',kind='trace',
        sources=['docs/evaluation/langsmith_check.json','docs/evaluation/agent_integration.json','docs/evaluation/tests-windows.txt'])
    add('성과는 모델 이름보다, 다음 판단까지 연결한 흐름입니다','교훈 · 트레이드오프 · 다음 개선',[
        '구현: 데이터 점검 → 계획 승인 → 실제 비교 → 다음 실험 제안',
        '교훈: Recall을 높이면 정상 오판도 늘 수 있고, 선택 규칙이 결론을 바꿉니다.',
        '한계: 합성 데이터·작은 불량 표본·반복 Validation·Test 재사용 탐색',
        '다음: 현장 기준 합의 → 독립 기간 검증 → 엔지니어 작업시간 측정'],25,
        '핵심 성과는 가장 좋은 모델 이름을 말하는 것보다, 계획 승인부터 비교 결과와 다음 실험까지 연결했다는 점입니다. Recall을 높이면 정상 오판이 늘 수 있고, 선택 기준이 결과를 바꾼다는 것을 배웠습니다. 합성 데이터와 작은 표본, 반복 비교가 한계입니다. 다음에는 현장 기준과 독립 기간 데이터로 검증하고 실제 작업시간 절감도 측정하겠습니다. 감사합니다.')
    full=[['모델 · 전체 제품 학습','불균형 설정','Recall','정상 오판률','F1']]+[
        [NAMES[x['model']],x['balance'],pct(x['validation']['recall']),pct(x['validation']['fpr']),pct(x['validation']['f1'])] for x in ranks]
    add('모델별 Validation 비교','부록 · 알고리즘마다 최선 불균형 설정 1개 표시',[
        '같은 전체 제품 Train에서 비교했습니다. 이 표는 제품 조합 확장 전 단계입니다.',
        '후보별 판정값은 Validation에서 조정했습니다. Test 전체 모델 순위가 아닙니다.'],table=full,
        note='전체 제품 학습 24조건에서 모델마다 최선 불균형 설정을 하나씩 표시합니다. CatBoost와 HistGradientBoosting의 제품 조합을 확장해 최종 CatBoost A+B+C를 선택했습니다. 모든 후보가 두 목표를 못 맞추면 목표 부족분 합계가 적은 탐색 참고 후보를 보여 줍니다.',sources=['docs/evaluation/synthetic_process.json'])
    other=[['데이터','기본 / 선택·참고 후보','불량 발견','정상 오판','판단']]
    for name,label in [('synthetic_process','합성 센서'),('secom','UCI SECOM'),('images','합성 이미지')]:
        q=receipts[name]['result']; a=q['baseline_test']; z=q['test']
        other.append([label,('축소 ResNet' if name=='images' else 'Logistic')+' / '+NAMES[q['winner']['model']],
                      f'{a["tp"]} → {z["tp"]} / {z["positive_count"]}',f'{a["fp"]} → {z["fp"]} / {z["tn"]+z["fp"]}',
                      '목표 충족 후보 없음' if name=='secom' else '불량 표본 보강' if name=='images' else '추가 실험 필요'])
    add('데이터를 바꾸면 후보와 결론도 달라집니다','부록 · 각 행 안에서만 같은 Test 비교',[
        '합성 센서·이미지는 실제 공장 데이터가 아닙니다. SECOM은 실제 공개 센서 자료입니다.',
        '각 데이터의 Test는 서로 다릅니다. 모델의 종합 순위로 해석하지 않습니다.'],table=other,
        note='SECOM은 1,567건 중 불합격104건이고 제조사·제품ID는 제공되지 않습니다. 이미지 Test의 불량15건은 작은 표본이어서 목표 통과를 확정하기 어렵습니다.',
        sources=['https://archive.ics.uci.edu/dataset/179/secom','https://doi.org/10.24432/C54305','docs/evaluation/summary.json'])
    before,after=retrain['before'],retrain['after']
    add('3σ는 변화 감지, Recall 목표는 판정 성능입니다','부록 · 재학습 판단과 기존 모델 관리',[
        '3σ: 공정 값·오차가 평소와 다른가? 안정성·분포·관리도 가정을 먼저 확인',
        'Recall/FPR: 정답이 있는 새 데이터에서 검사 모델이 목표를 만족하는가?',
        '현재 구현은 교육용 Recall 90%·정상 오판 15% 점검 · 3σ 트리거 미구현',
        f'새 합성 배치의 동일 Test: 불량 발견 {before["tp"]} → {after["tp"]} / {after["positive_count"]}개',
        f'정상 오판 {before["fp"]} → {after["fp"]}개 · 승인 후 같은 알고리즘 재학습'],
        note='정답이 없으면 Recall을 직접 계산할 수 없습니다. 향후 드리프트 감지를 검토할 수 있지만 변화 감지만으로 재학습 효과를 확정하지 않습니다. 현재는 새 정답 데이터의 Train 구간으로 점검하고 승인 후 후보를 비교합니다. 운영 모델을 자동 교체하지 않습니다.',sources=['app/monitoring.py','docs/evaluation/model_management.json'])
    add('6:2:2는 사전 고정, 혼합 비율은 별도 실험입니다','부록 · 데이터 분할과 앙상블',[
        'Train: 모델 학습 / Validation: 모델·제품 조합·판정값 선택 / Test: 선택 결과 확인',
        '합성 센서 768 / 256 / 256 · 후보 A+B+C는 Train 중 544건 학습',
        'SECOM은 날짜 블록 60/20/20 → 실제 행 수 820 / 367 / 380',
        'RF·부스팅은 사용 · 서로 다른 모델의 혼합 가중치 탐색은 미구현',
        '확장 시: 사전에 정한 가중치를 Validation으로 고르고 새 Test로 확인'],
        note='6대2대2가 모든 데이터의 최적 비율이라는 뜻은 아닙니다. 수업 기준으로 정한 재현 가능한 실험 조건입니다. 같은 lot·날짜 경계를 지키므로 실제 행 비율은 달라집니다. 추가 탐색으로 기존 Test를 반복 확인했다면 독립 기간 Test를 새로 확보해야 합니다.',sources=['app/data.py','docs/MODEL_DISCOVERY_DESIGN.md'])
    add('구현한 기술과 남은 범위를 구분합니다','부록 · 최종 보고서 기술 대응표',[
        '과제의 기술 목록 전체를 구현했다고 주장하지 않습니다.'],table=[
        ['요소','구현·증빙','현재 한계'],
        ['Multi-Agent / LangGraph','점검·검색·계획·결과 검토 역할','고정 그래프 · 자율 Supervisor 미구현'],
        ['RAG','TF-IDF / 선택적 Qwen 임베딩·cosine','별도 Vector DB 미도입 · 14개 소규모 근거'],
        ['외부 모델 탐색','공식 문서 검색·ExtraTrees/GaussianNB','검토된 두 adapter만 지원'],
        ['MCP','stdio 4도구 + 지식 Resource','승인·학습 실행은 API에서만'],
        ['코드 품질','모듈화·DI·단위/API 테스트','배포·동시 다중 사용자 미검증'],
        ['LangSmith','실제 root/child trace API 재조회','LLM 응답 성공과 별도'],
        ['LLM-as-a-Judge','미구현','허용 코드·출처 검사를 대신 사용']],
        note='기존 dense/TF-IDF 검색 평가는 고정13질의에서 Hit@3 양쪽100%, MRR@3 양쪽0.9615였습니다. 일반적인 검색 성능 우위를 주장하지 않습니다. LLM을 꺼도 규칙 실험 모드는 동작하며 실행 상태를 UI와 보고서에 표시합니다.',sources=['docs/ASSIGNMENT_CHECKLIST.md','docs/evaluation/agent_integration.json'])
    add('출처와 재현 자료','부록 · 질문 대비와 제출 링크',[
        '모델 참고: 스칼라 9주차 수업 + 15주차 반도체 발표자료',
        'SECOM: McCann & Johnston (2008), UCI · DOI 10.24432/C54305 · CC BY 4.0',
        '측정 원본: docs/evaluation/ · 고정 seed 45 · 기존 Mac 결과와 Windows 검증 구분',
        '소스·최종 보고서: github.com/dragonchan95/Semiconductor-AI-Agent',
        '기존 Google Slides 주소는 README에 보존 · 온라인 갱신 여부는 별도 확인'],
        note='발표는 1~10장이고 11~16장은 질문용입니다. 회사 공통 Recall 수치나 수율·금액 이익, 시간 절감은 확인하지 않았습니다. 제출 명단 엑셀과 교육자료 원본, API 키는 공개 저장소에 포함하지 않습니다.',sources=[REPO_URL,SLIDES_URL,'docs/CURRICULUM_MODELS.md'])
    assert sum(s['seconds'] for s in content)==280
    return content,receipts,trace


def color(value):
    return RGBColor.from_string(value)


def box(slide,x,y,w,h,fill,line=None,rounded=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=color(fill)
    if line:shape.line.color.rgb=color(line)
    else:shape.line.fill.background()
    if rounded:
        shape.adjustments[0]=.08
    return shape


def text(slide,value,x,y,w,h,size=20,fill=INK,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=shape.text_frame;frame.word_wrap=True
    frame.margin_left=frame.margin_right=Inches(.02)
    frame.margin_top=frame.margin_bottom=Inches(.02)
    for i,line in enumerate(str(value).split('\n')):
        p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
        p.text=line;p.font.name='맑은 고딕';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fill)
        p.space_after=Pt(8);p.line_spacing=1.14
    return shape


def add_table(slide,rows,y=3.05,height=3.25,font=17,widths=None):
    cols=len(rows[0]);shape=slide.shapes.add_table(len(rows),cols,Inches(.65),Inches(y),Inches(12.03),Inches(height))
    table=shape.table
    if widths:
        for i,v in enumerate(widths):table.columns[i].width=Inches(v)
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            cell=table.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.13);cell.margin_right=Inches(.1)
            cell.margin_top=Inches(.07);cell.margin_bottom=Inches(.06);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=color(TEAL if i==0 else ('FFFFFF' if i%2 else PALE))
            for p in cell.text_frame.paragraphs:
                p.font.name='맑은 고딕';p.font.size=Pt(font);p.font.bold=i==0;p.font.color.rgb=color('FFFFFF' if i==0 else INK)
    return shape


def picture(slide,path,x,y,w,h):
    from PIL import Image
    iw,ih=Image.open(path).size
    scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
    box(slide,x-.08,y-.08,w+.16,h+.16,'FFFFFF',LINE,True)
    slide.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))


def build_deck(content,trace):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    for i,item in enumerate(content,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=color(BG)
        box(slide,0,0,.16,7.5,TEAL)
        text(slide,item['context'],.65,.32,12,.3,12,TEAL,True)
        text(slide,item['title'],.65,.91,12,1.1,32 if '\n' not in item['title'] else 38,INK,True)
        box(slide,.65,6.91,12.03,.015,LINE)
        text(slide,'SKALA 광주 2반 · G045 김용찬',.65,7.02,8,.25,10,MUTED)
        text(slide,f'{i:02d} / {len(content):02d}'+('  부록' if i>10 else ''),11.1,7.02,1.6,.25,10,MUTED)
        slide.notes_slide.notes_text_frame.text=item['note']+'\n\n출처: '+'; '.join(item['sources'])
        if item['kind']=='cover':
            text(slide,item['body'][0],.72,2.75,11.5,.8,23,MUTED)
            text(slide,item['body'][1],.72,3.77,11.5,.5,20,INK)
            for j,(a,b) in enumerate([('01','데이터 점검'),('02','계획 승인'),('03','실제 비교'),('04','다음 실험')]):
                x=.72+j*3.08;box(slide,x,5.15,2.78,1.05,PALE,rounded=True)
                text(slide,a,x+.2,5.27,.5,.32,14,TEAL,True);text(slide,b,x+.2,5.68,2.4,.4,19,INK,True)
        elif item['kind']=='architecture':
            text(slide,item['body'][0],.68,2.12,12,.5,20,MUTED)
            nodes=[('데이터 점검','행·결측·라벨·분할'),('근거 검색','교육 근거 / 공식 문서'),('실험 계획','모델·전처리·예산'),('엔지니어 승인','계획·데이터 hash 확인')]
            for j,(a,b) in enumerate(nodes):
                x=.7+j*3.08;box(slide,x,3.02,2.77,1.26,PALE,rounded=True)
                text(slide,a,x+.15,3.2,2.48,.4,21,TEAL,True);text(slide,b,x+.15,3.77,2.5,.35,13,MUTED)
                if j<3:text(slide,'→',x+2.8,3.38,.28,.4,19,TEAL,True)
            box(slide,.7,4.8,5.7,1.35,'FFFFFF',LINE,True)
            text(slide,'승인 후 학습·평가',.95,5,5.15,.4,22,TEAL,True)
            text(slide,'Train 학습 → Validation 선택 → 고정 Test 확인',.95,5.57,5.15,.35,15,MUTED)
            box(slide,6.85,4.8,5.7,1.35,ORANGE,rounded=True)
            text(slide,'결과 검토 → 다음 실험',7.1,5,5.15,.4,22,INK,True)
            text(slide,'새 후보 탐색 / 기존 모델 점검 → 새 승인',7.1,5.57,5.15,.35,16,MUTED)
            text(slide,'MCP는 조회·계획 도구를 제공하고, LangSmith는 실행 경로를 추적합니다.',.75,6.4,11.8,.3,13,MUTED)
        elif item['image']:
            for j,line in enumerate(item['body']):
                text(slide,line,.72,2.25+j*.84,5.75,.74,21 if j==0 else 18,INK if j==0 else MUTED,j==0)
            picture(slide,DOCS/'assets'/item['image'],7.1,2.05,5.3,4.55)
        elif item['kind']=='trace':
            for j,line in enumerate(item['body']):text(slide,line,.7,2.15+j*.62,12,.52,19,INK,j==2)
            box(slide,.72,4.93,11.9,1.42,'FFFFFF',LINE,True)
            text(slide,'experiment.propose → LangGraph → 점검 · 검색 · 계획',.97,5.1,11.3,.5,23,TEAL,True)
            if trace:text(slide,'API 재조회 성공 · child run '+str(len(trace['child_runs']))+'개 · '+trace['root_run_id'],.97,5.86,11.1,.3,13,MUTED)
            text(slide,'현재 Windows: LLM 미사용 / 기존 Mac: 역할별 verified·fallback 기록 보존 · LLM-as-a-Judge 미구현',.75,6.5,11.8,.3,12,MUTED)
        elif item['table']:
            body_size=18 if i!=11 else 15
            for j,line in enumerate(item['body']):text(slide,line,.7,2.02+j*.46,12,.44,body_size,MUTED)
            y=2.12+len(item['body'])*.46+.16
            height=6.64-y
            font=16 if i<=10 else 13
            if i==11:font=12
            widths=None
            if i in [6,11]:widths=[5.43]+([2.2]*3) if i==6 else [4.43,1.9,1.9,1.9,1.9]
            if i==7:widths=[5.23,3.4,3.4]
            if i==15:widths=[2.35,5,4.68]
            add_table(slide,item['table'],y,height,font,widths)
        else:
            for j,line in enumerate(item['body']):
                y=2.16+j*.85;box(slide,.7,y,.055,.56,TEAL)
                text(slide,line,.95,y,11.6,.7,21 if i<=10 else 19,INK)
    prs.core_properties.title='반도체 검사 모델 실험 도우미'
    prs.core_properties.author='SKALA 광주 2반 · G045 김용찬'
    target=DOCS/'G045_Semiconductor_Agent.pptx';prs.save(target)
    return target


def write_documents(content,receipts,trace):
    (DOCS/'slide-content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
    script='SKALA 광주 2반 · G045 김용찬 / 과제 2\n반도체 검사 모델 실험 도우미\n'
    script+='2026-10-22 발표 · 본문 10장 + 질문용 부록 6장\n발화 280초 + 화면 전환 20초 = 5분 계획\n\n'
    elapsed=0
    report='# 반도체 검사 모델 실험 도우미\n\nSKALA 광주 2반 · G045 김용찬 · 과제 2 · 2026-10-22 발표\n\n'
    report+='새 검사 데이터에 맞는 모델을 비교하고 결과에 따라 다음 실험을 제안하는 Agent입니다.\n\n'
    report+='[발표 PPTX](G045_Semiconductor_Agent.pptx) · [발표 PDF](G045_Semiconductor_Agent.pdf) · [5분 대본](발표대본_G045.txt)\n\n'
    report+='원래 Mac에서 측정한 모델 실험 기록을 보존하고, Windows의 테스트·추적 검증을 별도로 남겼습니다.\n\n'
    for i,s in enumerate(content,1):
        elapsed+=s['seconds'];timing=f'{s["seconds"]}초 / 누적 {elapsed//60}:{elapsed%60:02d}' if i<=10 else '질문용 부록'
        script+=f'[{i}장 / {timing}] '+s['title'].replace('\n',' ')+'\n'+s['note']+'\n\n'
        report+=f'## {i}. '+s['title'].replace('\n',' ')+'\n\n'+s['context']+'\n\n'
        report+='\n\n'.join(s['body'])+'\n\n'
        if s['table']:
            report+='| '+' | '.join(s['table'][0])+' |\n|'+'---|'*len(s['table'][0])+'\n'
            report+='\n'.join('| '+' | '.join(str(c).replace('\n',' / ') for c in row)+' |' for row in s['table'][1:])+'\n\n'
        if s['image']:report+=f'![실제 앱 화면](assets/{s["image"]})\n\n'
        report+=s['note']+'\n\n'
        if s['sources']:report+='출처·증빙: '+' · '.join(s['sources'])+'\n\n'
    report+='''## 아키텍처 흐름도

```mermaid
flowchart LR
    A[센서 CSV / 이미지 manifest] --> B[데이터 점검]
    B --> C[근거 검색 RAG]
    C --> D[실험 계획 제안]
    D --> E{엔지니어 승인}
    E -->|승인| F[Train 실제 학습]
    E -->|반려| D
    F --> G[Validation 후보·판정값 선택]
    G --> H[고정 Test 확인]
    H --> I[결과 검토]
    I --> J[새 모델 탐색 / 기존 모델 점검]
    J --> D
    K[MCP 조회·계획 도구] --> D
    L[LangSmith 추적] -.-> B
    L -.-> D
    L -.-> I
```

승인은 SQLite에 저장한 계획 hash와 데이터 hash로 검증합니다. LangGraph interrupt를 사용한 구현은 아닙니다.
외부 모델 탐색 후 승인과 실행은 API에서 처리합니다. 운영 모델 자동 교체는 하지 않습니다.

## 수행하며 배운 점

LLM 설명과 실험의 수치·실행 권한을 분리해야 잘못된 수치나 무승인 학습을 막을 수 있었습니다.
목표와 분할을 사전에 정하지 않으면 높은 F1, 높은 Recall 중 어떤 결론을 택할지 임의로 바뀝니다.
Mac에서 통과하던 SQLite 코드도 Windows에서는 연결이 닫히지 않아 파일을 잠글 수 있었습니다.
환경 이동에서는 소스 코드뿐 아니라 설정, 모델 설치 여부, 실행 이력을 구분해 복원해야 합니다.

## 한계와 후속 평가

단일 seed·동일 Validation의 반복 선택·작은 불량 표본이 한계입니다.
새 모델 탐색에 이미 확인한 Test를 재사용했으므로 독립 기간의 새 Test가 필요합니다.
전처리·모델·판정값·학습 제품 범위가 함께 바뀐 비교이며 특정 기법의 인과 효과가 아닙니다.
수율, 금액 이익, 작업시간 절감, 양산 효과는 측정하지 않았습니다.
회사 공통 Recall/FPR 숫자를 가정하지 않고 현장 담당자가 목적과 비용에 맞춰 정해야 합니다.
LLM-as-a-Judge·별도 Vector DB·자율 Supervisor·3σ 재학습 트리거·혼합 가중치 탐색은 미구현입니다.
'''
    script+='''질문 대비
Q. AdaBoost는 어디서 나온 모델인가요?
A. 15주차 반도체 발표자료 40~42쪽의 Adaboost Regression / Ada Boost 회귀 사례를 참고해 분류용 AdaBoostClassifier로 적용했습니다. 새 모델 탐색 RAG가 추가한 모델은 ExtraTrees와 GaussianNB입니다.

Q. 왜 F1이 더 높은 AdaBoost를 선택하지 않았나요?
A. 이번 실험은 사전에 정한 Recall 목표와 정상 오판 한도 충족을 먼저 보고, Recall 우선·동률이면 정상 오판이 낮은 후보를 골랐습니다. 다른 업무 목표라면 선택 규칙도 바뀔 수 있습니다.

Q. 6:2:2가 항상 최적인가요?
A. 수업 기준으로 정한 사전 고정 조건입니다. 모든 데이터에서 최적이라는 뜻은 아닙니다. 같은 lot·날짜 경계를 보존하고 Test를 보고 비율을 고르지 않습니다.

Q. 자동화와 Agent는 어떤 차이가 있나요?
A. 현재는 역할을 나눈 제한된 Agent 구조입니다. 근거를 검색해 설명하고 계획을 제안하며 도구를 실행한 결과로 다음 실험을 연결합니다. 실제 수치·후보 순위·승인은 규칙 엔진이 책임집니다. 자율 Supervisor나 무제한 모델 발굴을 구현했다고 주장하지 않습니다.

Q. LangSmith가 연결되면 LLM도 작동하나요?
A. 추적 서버와 LLM 추론 서버는 별도입니다. 현재 Windows에서는 추적을 확인했고 로컬 LLM은 미사용입니다. 기존 Mac의 LLM verified/fallback 기록을 따로 보존했습니다.

Windows 앱 실행: continued-project 폴더의 start.cmd를 더블클릭하고 http://localhost:8700을 엽니다.
새 환경은 예제 데이터 준비 → 계획 제안 → 승인·학습 순서입니다. 발표는 기록된 화면을 설명하므로 새 학습을 기다릴 필요가 없습니다.
'''
    (DOCS/'REPORT.md').write_text(report,encoding='utf-8')
    (DOCS/'발표대본_G045.txt').write_text(script,encoding='utf-8')
    r=receipts['synthetic_process']['result'];b=r['baseline_test'];t=r['test']
    readme=f'''# 반도체 검사 모델 실험 도우미

**SKALA 광주 2반 · G045 김용찬 · 신규 AI Agent 과제 2**

검사 데이터를 점검하고, 승인받아 모델을 비교한 뒤, 결과에 따라 다음 실험을 제안합니다.

[최종 보고서](docs/REPORT.md) · [발표 PPTX](docs/G045_Semiconductor_Agent.pptx) · [발표 PDF](docs/G045_Semiconductor_Agent.pdf) · [5분 대본](docs/발표대본_G045.txt) · [과제 대응표](docs/ASSIGNMENT_CHECKLIST.md)

기존 [Google Slides]({SLIDES_URL}) 주소를 보존했습니다. 이번 수정의 원본은 이 저장소의 보고서와 PPTX입니다. Slides의 온라인 갱신 상태는 별도로 확인해야 합니다.

## 하나의 사용 흐름

데이터 선택 → 점검·근거 검색 → 계획 확인·승인 → 실제 학습 → Validation 비교 → 고정 Test 확인 → 새 모델 탐색 / 기존 모델 점검

![같은 Test의 기본·후보 비교](docs/assets/app-result.jpg)

## Windows에서 실행

`start.cmd`를 더블클릭하고 **http://localhost:8700**을 엽니다. 실행 창은 열어 둡니다.
새 설치는 PowerShell에서 `./start.ps1`로 Python 환경과 의존성을 준비합니다. Python 3.11 이상이 필요합니다.
이 작업 폴더에서는 상위 폴더의 `.venv`도 인식합니다. 처음에는 **예제 데이터 준비**를 누릅니다.
새 환경에는 과거 모델 파일과 DB가 없으므로 계획 승인 후 학습해야 새 실행 이력이 생깁니다.
저장된 실측 결과는 `docs/evaluation/`에서 확인할 수 있습니다.

macOS/Linux에서는 `bash start.sh`를 실행합니다. macOS의 XGBoost·LightGBM은 OpenMP 설치가 필요할 수 있습니다.

## 현재 모드와 LangSmith

`.env.example`을 참고해 로컬 `.env`를 설정합니다. 키는 Git에서 제외합니다.
APAC LangSmith에서 **실제 계획 생성과 하위 실행 기록의 API 재조회**를 확인했습니다.
[추적 증빙](docs/evaluation/langsmith_check.json) · [설정 안내](docs/LANGSMITH_SETUP.md)

현재 Windows 모드는 LLM 미사용·TF-IDF 검색·규칙 설명입니다. Ollama와 Qwen 모델을 설치하고
`SEMI_LLM_ENABLED=1`로 설정하면 로컬 LLM·임베딩을 사용합니다.
기존 Mac의 LLM 검증 기록은 [agent_integration.json](docs/evaluation/agent_integration.json)에 보존합니다.
LangSmith 연결과 LLM 응답 성공은 별개이며 UI에 verified/fallback/disabled 상태를 표시합니다.

## 고정한 조건과 결과

6:2:2를 사전에 고정했습니다. 합성 센서는 Train768 / Validation256 / Test256이며 같은 lot을 섞지 않습니다.
SECOM은 날짜 블록60/20/20이라 실제 행 수는 820 / 367 / 380입니다.
여러 모델과 불균형 설정을 비교한 뒤 상위 후보의 제품 조합을 확장합니다. 최대52조건이며 완전탐색이 아닙니다.
후보는 Validation에서 고르고 Test로 순위를 바꾸지 않습니다.

|같은 합성 센서 Test|기본 Logistic|선택 후보 CatBoost A+B+C|
|---|---:|---:|
|불량 발견|{b['tp']}/32|{t['tp']}/32|
|정상 오판|{b['fp']}/224|{t['fp']}/224|
|Recall|{pct(b['recall'])}|{pct(t['recall'])}|

후보 Recall이 교육용 목표90%에 못 미쳐 추가 실험이 필요합니다. 정상 오판 목표15%도 함께 확인합니다.
AdaBoost는 F1이 더 높지만 이번 Recall 우선 선택 규칙에서는 CatBoost가 앞섰습니다.
모델별 비교와 세 데이터의 결과는 [보고서](docs/REPORT.md)에 있습니다.
측정은 기존 Mac 실험 기록을 보존하며 Windows의 재검증은 별도 증빙으로 구분합니다.

## 구현 범위

LangGraph 역할 분리, 교육 근거14개 검색과 공식 문서 기반 새 모델 탐색, MCP4도구·Resource,
의존성 주입, 승인 hash, 실제 모델 학습·파일 저장, 모델 점검·재학습, 제품별 학습 비교를 구현했습니다.
외부 모델 탐색은 **ExtraTrees·GaussianNB의 검토된 adapter**를 지원합니다.
현재 RAG는 TF-IDF 또는 Qwen 임베딩·cosine 검색이며 별도 Vector DB는 없습니다.
자율 Supervisor·LLM-as-a-Judge·3σ 트리거·임의 코드 생성·앙상블 혼합 가중치 탐색은 미구현입니다.

운영 모델은 자동 교체하지 않습니다. 회사 공통 Recall 수치, 양산 수율·금액 이득·작업시간 절감은 검증하지 않았습니다.
작은 불량 표본·반복 Validation과 추가 탐색의 Test 재사용이 한계이며 독립 기간 검증이 필요합니다.
교육 원본·제출 명단 엑셀·비밀 값은 공개 저장소에 넣지 않습니다.

## 검증과 재현

```powershell
$env:LANGSMITH_TRACING='false'
..\\.venv\\Scripts\\python.exe -X utf8 -m unittest discover -s tests -v
..\\.venv\\Scripts\\python.exe -X utf8 scripts/check_mcp.py
..\\.venv\\Scripts\\python.exe -X utf8 scripts/check_langsmith.py
..\\.venv\\Scripts\\python.exe -X utf8 scripts/build_report.py
```

독립 설치에서는 `..\\.venv` 대신 `.venv`를 사용합니다. 발표 생성 패키지는 `requirements-presentation.txt`에 있습니다.
Windows에서는 SQLite 연결을 transaction 종료 후 명시적으로 닫아 DB 파일 잠금 문제를 수정했습니다.
SECOM: McCann & Johnston(2008), [UCI 원본](https://archive.ics.uci.edu/dataset/179/secom), DOI10.24432/C54305, CC BY4.0.
센서 원파일590열과 UCI 메타데이터591 features 표기를 구분합니다. 제조사·제품ID는 원자료에 제공되지 않습니다.
'''
    (ROOT/'README.md').write_text(readme,encoding='utf-8')


def main():
    content,receipts,trace=make_content()
    target=build_deck(content,trace)
    write_documents(content,receipts,trace)
    print(json.dumps({'pptx':str(target),'slides':len(content),'main_slides':10,
                      'planned_speech_seconds':sum(s['seconds'] for s in content)},ensure_ascii=False))


if __name__=='__main__':
    main()
