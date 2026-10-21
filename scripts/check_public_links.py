"""Verify public, unauthenticated access to the repository and existing Slides export."""
import json
from pathlib import Path
import httpx

ROOT=Path(__file__).resolve().parent.parent
urls={
    'github_report':'https://raw.githubusercontent.com/dragonchan95/Semiconductor-AI-Agent/main/docs/REPORT.md',
    'google_slides_export':'https://docs.google.com/presentation/d/1J2m_cpXyvb-5_Xabs2DsD3EnEV6B6GicF5OP2s5FohU/export/pptx'}
receipt={}
with httpx.Client(follow_redirects=True,timeout=45) as client:
    for name,url in urls.items():
        response=client.get(url)
        valid=response.status_code==200 and ('# 반도체 검사 모델 실험 도우미' in response.text if name=='github_report' else response.content.startswith(b'PK'))
        receipt[name]={'http_status':response.status_code,'unauthenticated_access_verified':valid,
                       'content_type':response.headers.get('content-type')}
        if valid and name=='google_slides_export':
            target=ROOT/'var/public-slides-before-update.pptx';target.write_bytes(response.content)
            from pptx import Presentation
            prs=Presentation(target)
            receipt[name]['slides']=len(prs.slides)
            receipt[name]['current_slide_titles']=[next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()),'') for slide in prs.slides]
receipt['google_slides_updated_this_session']=False
(ROOT/'docs/evaluation/public_links.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
