"""Stage the requested workbook; leave every other student's cells unchanged."""
from pathlib import Path
import json
from copy import copy
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment

ROOT=Path(__file__).resolve().parent.parent
WEEK=ROOT.parent.parent
SOURCE=WEEK/'AI Agent Mini-Project-Semiconductor-AI-Agent.xlsx'
TARGET=ROOT/'docs/submission'/SOURCE.name
TARGET.parent.mkdir(parents=True,exist_ok=True)
wb=load_workbook(SOURCE)
assert wb['2반']['A11'].value=='G045' and wb['2반']['B11'].value=='김용찬'
before={(s.title,c.coordinate):c.value for s in wb for row in s for c in row if c.value is not None}
ws=wb['2반']
ws['C11']=2
ws['D11']='\n'.join([
    '반도체 검사 모델 실험 도우미 (신규 AI Agent · 과제 2)',
    '- 데이터 점검 → 근거 검색 → 계획 승인 → 실제 학습 → 결과·다음 실험',
    '- 스칼라 수업·반도체 자료를 참고한 후보 모델 비교',
    '- 6:2:2 사전 고정·같은 lot/날짜 분리·Validation 후보 선택',
    '- Recall 우선·정상 오판·Precision·F1 함께 비교',
    '- 공식 문서 탐색 RAG·ExtraTrees/GaussianNB 승인 비교·목록 갱신',
    '- 기존 모델 점검·같은 알고리즘 재학습·제품별 학습 비교',
    '- LangGraph·MCP 4도구/Resource·DI·모듈화·30개 단위/API 테스트',
    '- LangSmith APAC 실제 계획·하위 실행 API 재조회 확인',
    '- Windows 실행 복구·5분 발표 10장 + 질문용 부록 6장',
    '- 현재 Windows LLM 미사용·기존 Mac의 LLM 검증 기록 별도 보존',
    '- 양산 효과·자동 배포·LLM-as-a-Judge 미구현'])
repo='https://github.com/dragonchan95/Semiconductor-AI-Agent'
report=repo+'/blob/main/docs/REPORT.md'
slides='https://docs.google.com/presentation/d/1J2m_cpXyvb-5_Xabs2DsD3EnEV6B6GicF5OP2s5FohU/edit'
ws['E11']=repo
ws['F11']=slides
for addr in ['E11','F11']:
    ws[addr].hyperlink=ws[addr].value
    font=copy(ws[addr].font);font.color='0563C1';font.underline='single';ws[addr].font=font
ws['F11'].comment=Comment('요청한 기존 Google Slides 주소를 유지합니다. 최신 PPTX는 준비되었으며 온라인 원본 반영·링크 전체 공유는 별도 확인이 필요합니다. 최신 로컬 내용의 보고서: '+report,'G045')
name='G045 제출 안내'
if name in wb.sheetnames:
    info=wb[name]
else:info=wb.create_sheet(name)
rows=[['항목','내용'],['소속','SKALA 광주 2반 · G045 김용찬 · 과제 2'],
      ['최종 보고서 Google Slides URL',slides],['소스 코드 URL',repo],['최신 보고서 사본',report],
      ['최신 발표 PPTX',repo+'/blob/main/docs/G045_Semiconductor_Agent.pptx'],
      ['최신 발표 PDF',repo+'/blob/main/docs/G045_Semiconductor_Agent.pdf'],
      ['5분 대본',repo+'/blob/main/docs/발표대본_G045.txt'],
      ['온라인 Slides 상태','이 세션에서 온라인 갱신 미확인. 최신 로컬 PPTX/PDF와 GitHub 보고서를 사용.'],
      ['제출 일정','2026-10-22(목) 13:00 마감 / 13:10 발표 시작 / 개인 5분'],
      ['제출 절차','이 파일은 로컬 명단입니다. 지정된 공용 Google Spreadsheet 등록 여부는 별도 확인 필요.'],
      ['LangSmith','APAC에서 실제 계획 생성·하위 실행 API 재조회 성공. 키는 로컬 .env에만 저장.'],
      ['실행 모드','현재 Windows는 LLM 미사용. 기존 Mac LLM verified/fallback 기록과 구분.'],
      ['발표 구성','본문 10장 280초 + 전환 20초 / 질문용 부록 6장']]
for i,row in enumerate(rows,1):
    for j,value in enumerate(row,1):
        cell=info.cell(i,j,value);cell.alignment=Alignment(vertical='top',wrap_text=True)
        cell.font=Font(name='맑은 고딕',size=11,color='173A49',bold=i==1)
        if value.startswith('https://'):cell.hyperlink=value;cell.font=Font(name='맑은 고딕',size=11,color='0563C1',underline='single')
        if i==1:cell.fill=PatternFill('solid',fgColor='DDEFEF')
    info.row_dimensions[i].height=38 if i>1 else 25
info.column_dimensions['A'].width=26;info.column_dimensions['B'].width=105
info.freeze_panes='B2'
wb.save(TARGET)
after=load_workbook(TARGET)
changed=[]
for key,value in before.items():
    sheet,coordinate=key
    if after[sheet][coordinate].value!=value:changed.append(key)
assert set(changed)<={('2반',x) for x in ['C11','D11','E11','F11']},changed
print(json.dumps({'staged_workbook':str(TARGET),'changed_existing_cells':changed,'added_sheet':name},ensure_ascii=False))
