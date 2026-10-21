"""Check slide structure and render a contact sheet from PowerPoint's PDF export."""
import json
from pathlib import Path
import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT=Path(__file__).resolve().parent.parent
docs=ROOT/'docs'
prs=Presentation(docs/'G045_Semiconductor_Agent.pptx')
pdf=pymupdf.open(docs/'G045_Semiconductor_Agent.pdf')
content=json.loads((docs/'slide-content.json').read_text(encoding='utf-8'))
assert len(prs.slides)==len(pdf)==len(content)==14
out=ROOT/'var/presentation-preview';out.mkdir(parents=True,exist_ok=True)
sheet=Image.new('RGB',(1920,1160),'#e3eaed')
overflow=[]
for i,page in enumerate(pdf):
    pix=page.get_pixmap(matrix=pymupdf.Matrix(1,1),alpha=False)
    im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
    im.thumbnail((472,265))
    sheet.paste(im,((i%4)*480+4,(i//4)*290+20))
    for s in prs.slides[i].shapes:
        if s.left < 0 or s.top < 0 or s.left+s.width>prs.slide_width+1000 or s.top+s.height>prs.slide_height+1000:
            overflow.append(i+1)
    if i in [2,3,4,5,6,7,8,10,13,14]:page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False).save(out/f'slide-{i+1:02d}.png')
    assert prs.slides[i].notes_slide.notes_text_frame.text.strip()
    assert '\ufffd' not in page.get_text()
assert not overflow,overflow
assert content[2]['kind']=='criteria'
assert content[5]['kind']=='ensemble_screen'
assert content[6]['kind']=='result_screen'
assert sum(x['seconds'] for x in content)==255
for i in [3,4,5,6,7]:
    assert any(s.shape_type==13 for s in prs.slides[i].shapes),f'Missing screenshot on slide {i+1}'
assert not any('Recall만 높으면' in x['title'] for x in content[10:])
assert not any('출처·시연' in x['title'] or x['title']=='구현 기술의 상세 범위' for x in content)
sheet.save(out/'contact-sheet.jpg')
receipt={'slides':14,'main_slides':10,'appendix_slides':4,'speaker_notes':14,
         'planned_speech_seconds':255,'planned_transition_seconds':20,'buffer_seconds':25,
         'timing_is_rehearsal_measurement':False,
         'criteria_before_results':True,'essential_ui_screens_slides':[4,5,6,7,8],
         'conversion':'Microsoft PowerPoint PDF export',
         'shape_bounds_ok':True,'no_replacement_characters':True,
         'source_metrics':'Preserved Mac 6:2:2 receipts; Windows verification stored separately',
         'google_slides_updated_this_session':False,'contact_sheet':'var/presentation-preview/contact-sheet.jpg'}
(docs/'evaluation/presentation_check.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False))
