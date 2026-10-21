"""Fresh unauthenticated GETs; never use Git credentials, cookies or API keys."""
from datetime import datetime,timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import subprocess

import httpx
from pptx import Presentation

ROOT=Path(__file__).resolve().parent.parent
repo='https://github.com/dragonchan95/Semiconductor-AI-Agent'
raw='https://raw.githubusercontent.com/dragonchan95/Semiconductor-AI-Agent/main/'
receipt={'verified_at':datetime.now(timezone.utc).isoformat(),'anonymous_requests':True,
         'commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'checks':{}}
with httpx.Client(follow_redirects=True,timeout=45,headers={'User-Agent':'G045-public-submission-check'}) as client:
    for label,local,path in [('report','docs/REPORT.md','docs/REPORT.md'),('pptx','docs/G045_Semiconductor_Agent.pptx','docs/G045_Semiconductor_Agent.pptx'),('pdf','docs/G045_Semiconductor_Agent.pdf','docs/G045_Semiconductor_Agent.pdf')]:
        r=client.get(raw+path+'?verify='+receipt['commit']);r.raise_for_status()
        local_bytes=(ROOT/local).read_bytes()
        exact=r.content==local_bytes
        # Git normalizes Markdown CRLF on Windows; binary artifacts stay exact.
        matched=(r.content.replace(b'\r\n',b'\n')==local_bytes.replace(b'\r\n',b'\n')) if label=='report' else exact
        receipt['checks'][label]={'url':raw+path,'http_status':r.status_code,'matches_latest_local_file':matched,
            'byte_identical':exact,'comparison':'UTF-8 Markdown with line endings normalized' if label=='report' else 'exact bytes',
            'sha256':sha256(r.content).hexdigest()}
        assert matched,f'{label} public file is stale'
    r=client.get(repo);r.raise_for_status()
    receipt['checks']['github_repository']={'url':repo,'http_status':r.status_code}
    r=client.get('https://docs.google.com/presentation/d/1J2m_cpXyvb-5_Xabs2DsD3EnEV6B6GicF5OP2s5FohU/export/pptx')
    google={'url':'https://docs.google.com/presentation/d/1J2m_cpXyvb-5_Xabs2DsD3EnEV6B6GicF5OP2s5FohU/edit','http_status':r.status_code}
    if r.status_code==200:
        deck=Presentation(BytesIO(r.content));google['slide_count']=len(deck.slides)
        titles=[' '.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
        google['matches_latest_story']=(len(deck.slides)==14
            and any('CatBoost 90%' in text and 'LightGBM 10%' in text for text in titles)
            and any('ABC 공통 모델' in text and 'D 전용 모델' in text for text in titles)
            and any('BGE-M3' in text and 'FAISS' in text for text in titles))
    else:google['matches_latest_story']=False
    receipt['checks']['google_slides']=google
receipt['google_slides_updated_this_session']=False
receipt['latest_report_and_artifacts_public']=all(receipt['checks'][x]['matches_latest_local_file'] for x in ['report','pptx','pdf'])
(ROOT/'docs/evaluation/public_links.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
