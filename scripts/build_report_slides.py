"""Put every page of the detailed report into a Google Slides upload PPTX.

This preserves the checked report layout, tables and diagrams exactly. Pages
are high-resolution images; full extracted text is retained in speaker notes.
The five-minute presentation deck is a separate artifact.
"""
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT/'docs'
PDF = DOCS/'G045_Final_Report.pdf'
TARGET = DOCS/'G045_Final_Report.pptx'


def main():
    pdf = pymupdf.open(PDF)
    deck = Presentation()
    deck.slide_width = Inches(pdf[0].rect.width/72)
    deck.slide_height = Inches(pdf[0].rect.height/72)
    full_text = []
    for page in pdf:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        pix = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), alpha=False)
        slide.shapes.add_picture(BytesIO(pix.tobytes('png')), 0, 0,
                               width=deck.slide_width, height=deck.slide_height)
        extracted = page.get_text()
        full_text.append(extracted)
        slide.notes_slide.notes_text_frame.text = '상세 최종 보고서 · 제출용 · '+str(page.number+1)+'/'+str(len(pdf))+'\n\n'+extracted
        # Preserve the report's clickable code/source/evidence links as hotspots.
        for link in page.get_links():
            if link.get('uri'):
                rect = link['from']
                target = slide.shapes.add_textbox(Inches(rect.x0/72), Inches(rect.y0/72),
                                                 Inches(rect.width/72), Inches(rect.height/72))
                target.click_action.hyperlink.address = link['uri']
    deck.core_properties.title = '반도체 검사 모델 실험 도우미 — 최종 보고서'
    deck.core_properties.author = 'SKALA 광주2반 G045 김용찬'
    deck.core_properties.subject = '교수님 제출용 상세 최종 보고서 / Google Slides 업로드용'
    deck.core_properties.comments = '각 슬라이드는 상세 보고서 PDF의 해당 쪽입니다. 본문 편집은 REPORT.md에서 하고 재생성합니다.'
    deck.save(TARGET)
    reopened = Presentation(TARGET)
    assert len(reopened.slides) == len(pdf) == 8
    assert all(sum(s.shape_type == 13 for s in slide.shapes) == 1 for slide in reopened.slides)
    assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in reopened.slides)
    text = ''.join(full_text)
    assert all(section in text for section in ['1. 표지', '2. 문제', '3. 적용', '4. 구현', '5. 교훈'])
    receipt = {'artifact': 'docs/G045_Final_Report.pptx', 'slides': len(reopened.slides),
               'source_pdf': 'docs/G045_Final_Report.pdf', 'source_sha256': sha256(PDF.read_bytes()).hexdigest(),
               'all_report_pages_included': True, 'required_sections': 5, 'render_dpi': 180,
               'page_size': 'A4 portrait', 'content_format': 'High-resolution page images; extracted text in notes; clickable PDF links preserved',
               'purpose': 'Upload this detailed report to Google Slides for the instructor; presentation deck is separate',
               'google_slides_online_update': False}
    (DOCS/'evaluation/report_slides_check.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(receipt, ensure_ascii=False))


if __name__ == '__main__':
    main()
